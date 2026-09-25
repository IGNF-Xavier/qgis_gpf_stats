import openpyxl

from geoplateforme_usage_stats.core.export_context import build_export_context
from geoplateforme_usage_stats.core.glossary import GLOSSARY
from geoplateforme_usage_stats.exporters import xlsx_exporter

EXPECTED_SHEETS = {
    "Dashboard", "Synthese", "Series_API", "Series_analytiques", "Series_groupes",
    "Controle_couverture", "Groupes", "Periode", "Journal_erreurs", "Glossaire",
}


def test_workbook_has_all_required_sheets(tmp_path, synthetic_dataset):
    results, groups, period = synthetic_dataset
    context = build_export_context(results, groups, period, errors=[])
    path = tmp_path / "report.xlsx"
    xlsx_exporter.export_workbook(str(path), context)

    wb = openpyxl.load_workbook(str(path))
    assert set(wb.sheetnames) == EXPECTED_SHEETS


def test_synthese_sheet_preserves_uuids(tmp_path, synthetic_dataset):
    results, groups, period = synthetic_dataset
    context = build_export_context(results, groups, period, errors=[])
    path = tmp_path / "report.xlsx"
    xlsx_exporter.export_workbook(str(path), context)

    wb = openpyxl.load_workbook(str(path))
    ws = wb["Synthese"]
    headers = [c.value for c in ws[1]]
    offering_id_col = headers.index("offering_id")
    values = [ws.cell(row, offering_id_col + 1).value for row in range(2, ws.max_row + 1)]
    assert "off-1" in values


def test_dashboard_sheet_has_kpi_rows_without_mixing_levels(tmp_path, synthetic_dataset):
    results, groups, period = synthetic_dataset
    context = build_export_context(results, groups, period, errors=[])
    path = tmp_path / "report.xlsx"
    xlsx_exporter.export_workbook(str(path), context)

    wb = openpyxl.load_workbook(str(path))
    ws = wb["Dashboard"]
    text_blob = "\n".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None)
    assert "Offerings" in text_blob
    assert "Endpoints" in text_blob
    assert "ne jamais additionner" in text_blob.lower() or "ne pas additionner" in text_blob.lower()


def test_glossary_sheet_matches_source(tmp_path, synthetic_dataset):
    results, groups, period = synthetic_dataset
    context = build_export_context(results, groups, period, errors=[])
    path = tmp_path / "report.xlsx"
    xlsx_exporter.export_workbook(str(path), context)

    wb = openpyxl.load_workbook(str(path))
    ws = wb["Glossaire"]
    rows = [(ws.cell(r, 1).value, ws.cell(r, 2).value) for r in range(2, ws.max_row + 1)]
    assert rows == GLOSSARY


def test_dashboard_omits_levels_not_in_the_selection(tmp_path, synthetic_dataset):
    results, groups, period = synthetic_dataset
    # keep only offerings - drop endpoint / producer_permission / consumer_permission / not_connected rows
    from geoplateforme_usage_stats.core.models import OFFERING

    offerings_only = [r for r in results if r.item.kind == OFFERING]
    context = build_export_context(offerings_only, [], period, errors=[])
    path = tmp_path / "report.xlsx"
    xlsx_exporter.export_workbook(str(path), context)

    wb = openpyxl.load_workbook(str(path))
    ws = wb["Dashboard"]
    text_blob = "\n".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None)
    assert "Offerings" in text_blob
    assert "Endpoints" not in text_blob
    assert "permissions producteur" not in text_blob.lower()
    assert "permissions consommateur" not in text_blob.lower()
    assert "Top 15 groupes" not in text_blob


def test_human_bytes_units():
    assert xlsx_exporter.human_bytes(500) == "500 o"
    assert xlsx_exporter.human_bytes(2048).endswith("Ko")
    assert xlsx_exporter.human_bytes(5 * 1024 * 1024).endswith("Mo")
    assert xlsx_exporter.human_bytes(5 * 1024 * 1024 * 1024).endswith("Go")


def test_workbook_survives_empty_dataset(tmp_path, synthetic_dataset):
    _, _, period = synthetic_dataset
    context = build_export_context([], [], period, errors=[])
    path = tmp_path / "empty.xlsx"
    xlsx_exporter.export_workbook(str(path), context)
    wb = openpyxl.load_workbook(str(path))
    assert set(wb.sheetnames) == EXPECTED_SHEETS
